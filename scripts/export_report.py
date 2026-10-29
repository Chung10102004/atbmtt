"""Export the local Vietnamese Markdown report to Word/PDF with embedded figures.

Install requirements-report.txt. This script reads aggregate evidence only and
does not load credentials, call external services, or run training.
"""
import os
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT/"reports/bao_cao_atbmhttt.md"


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            yield "code", "\n".join(code)
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",c) for c in row):
                    rows.append(row)
                i += 1
            yield "table", rows
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)",line)
            if match is None:
                raise ValueError("Unsupported figure syntax")
            yield "image",match.groups()
            i += 1
        elif line.startswith("#"):
            match = re.match(r"^(#+)\s+(.*)",line)
            yield "heading",(len(match[1]),match[2])
            i += 1
        elif line.startswith("- "):
            yield "bullet",line[2:]
            i += 1
        elif re.match(r"^\d+\.\s",line):
            yield "number",line
            i += 1
        else:
            paragraph = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(("#","|","- ","```","![")) and not re.match(r"^\d+\.\s",lines[i].strip()):
                paragraph.append(lines[i].strip())
                i += 1
            yield "paragraph"," ".join(paragraph)


def plain(text):
    return re.sub(r"\[([^]]+)\]\([^)]+\)",r"\1",text).replace("**","").replace("`","")


def pdf_inline(text):
    text = escape(text)
    text = re.sub(r"\*\*(.+?)\*\*",r"<b>\1</b>",text)
    text = re.sub(r"`(.+?)`",r"<font name='BodyFont'>\1</font>",text)
    return re.sub(r"\[([^]]+)\]\(([^)]+)\)",r'<link href="\2">\1</link>',text)


def figures():
    import pandas as pd
    os.environ.setdefault("MPLCONFIGDIR",str(ROOT/"reports/.mplcache"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter
    daily = pd.read_csv(ROOT/"reports/evidence/feature_study/daily_data_profile.csv")
    days = pd.to_datetime(daily.day,format="%Y/%m/%d")
    fig,axes = plt.subplots(2,1,figsize=(10,5.8),sharex=True,layout="constrained")
    axes[0].bar(days,daily.transactions,color="#416c93")
    axes[0].set_yscale("log")
    axes[0].set_ylabel("Transactions (log scale)")
    axes[1].plot(days,daily.positive_rate,marker="o",color="#ae6830")
    axes[1].yaxis.set_major_formatter(PercentFormatter(1))
    axes[1].set_ylabel("Laundering-label rate")
    for ax in axes:
        ax.axvline(pd.Timestamp("2022-09-11"),color="#a33d45",linestyle="--",label="11 Sep")
        ax.grid(axis="y",alpha=.2)
        ax.spines[["top","right"]].set_visible(False)
    axes[0].set_title("Observed daily distribution · IBM AML HI-Small")
    axes[1].set_xlabel("September 2022")
    axes[0].legend(frameon=False)
    for suffix in ("png","pdf"):
        fig.savefig(ROOT/f"reports/assets/daily_profile.{suffix}",dpi=180)
    plt.close(fig)


def word(items):
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    section = doc.sections[0]
    section.page_width,section.page_height = Cm(21),Cm(29.7)
    section.top_margin,section.bottom_margin = Cm(2.1),Cm(2)
    section.left_margin,section.right_margin = Cm(2.4),Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name,normal.font.size = "Arial",Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title","Heading 1","Heading 2"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor.from_string("214D69")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("ATBMHTTT | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)
    cover = True
    for kind,value in items:
        if kind == "heading":
            level,text = value
            if level == 1:
                p = doc.add_paragraph(plain(text),"Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(60)
            else:
                if cover:
                    doc.add_page_break()
                    cover = False
                doc.add_heading(plain(text),level=min(level-1,3))
        elif kind == "paragraph":
            p = doc.add_paragraph(plain(value))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cover else WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "bullet":
            doc.add_paragraph(plain(value),"List Bullet")
        elif kind == "number":
            p = doc.add_paragraph(plain(value))
            p.paragraph_format.space_after = Pt(2)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(value)
            run.font.name,run.font.size = "Consolas",Pt(8)
        elif kind == "table":
            table = doc.add_table(rows=0,cols=len(value[0]))
            table.style = "Table Grid"
            for i,row in enumerate(value):
                cells = table.add_row().cells
                for cell,content in zip(cells,row):
                    cell.text = plain(content)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.05
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.bold = i == 0
                    if i == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"),"E9F0F5")
                        cell._tc.get_or_add_tcPr().append(shade)
                if i == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
        elif kind == "image":
            caption,path = value
            doc.add_picture(str(REPORT.parent/path),width=Cm(16.2))
            p = doc.add_paragraph(plain(caption))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9)
                run.italic = True
    doc.core_properties.title = "Phát hiện giao dịch rửa tiền bằng học máy"
    doc.core_properties.subject = "Đồ án ATBMHTTT — đặc trưng thời gian và quan hệ tài khoản"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(REPORT.with_suffix(".docx"))


def pdf(items):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER,TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Table,TableStyle,Image,KeepTogether
    font_root = Path("C:/Windows/Fonts")
    fonts = {"BodyFont":"arial.ttf","BodyBold":"arialbd.ttf","BodyItalic":"ariali.ttf","BodyBoldItalic":"arialbi.ttf"}
    if not all((font_root/f).exists() for f in fonts.values()):
        import matplotlib
        font_root = Path(matplotlib.get_data_path())/"fonts/ttf"
        fonts = {"BodyFont":"DejaVuSans.ttf","BodyBold":"DejaVuSans-Bold.ttf",
                 "BodyItalic":"DejaVuSans-Oblique.ttf","BodyBoldItalic":"DejaVuSans-BoldOblique.ttf"}
    for name,path in fonts.items():
        pdfmetrics.registerFont(TTFont(name,str(font_root/path)))
    pdfmetrics.registerFontFamily("BodyFont",normal="BodyFont",bold="BodyBold",italic="BodyItalic",boldItalic="BodyBoldItalic")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ReportBody",fontName="BodyFont",fontSize=10.4,leading=14.8,spaceAfter=7,alignment=TA_JUSTIFY)
    head1 = ParagraphStyle("ReportH1",parent=body,fontName="BodyBold",fontSize=14,leading=19,spaceBefore=14,spaceAfter=9,keepWithNext=True,textColor=colors.HexColor("#214D69"))
    head2 = ParagraphStyle("ReportH2",parent=head1,fontSize=11.5,leading=16,spaceBefore=10)
    title = ParagraphStyle("ReportTitle",parent=head1,fontSize=20,leading=28,alignment=TA_CENTER)
    cell_style = ParagraphStyle("ReportCell",parent=body,fontSize=8,leading=11,spaceAfter=0,alignment=0)
    caption_style = ParagraphStyle("ReportCaption",parent=body,fontSize=8.8,leading=12,alignment=TA_CENTER)
    code_style = ParagraphStyle("ReportCode",parent=body,fontSize=8,leading=11,alignment=0,backColor=colors.HexColor("#F1F4F6"),borderPadding=6)
    document = SimpleDocTemplate(str(REPORT.with_suffix(".pdf")),pagesize=A4,
        rightMargin=2*cm,leftMargin=2.2*cm,topMargin=1.8*cm,bottomMargin=1.8*cm,
        title="Phát hiện giao dịch rửa tiền bằng học máy",author="",subject="Đồ án ATBMHTTT")
    story,cover = [],True
    for kind,value in items:
        if kind == "heading":
            level,text = value
            if level == 1:
                story.extend([Spacer(1,3*cm),Paragraph(pdf_inline(text),title),Spacer(1,.7*cm)])
            else:
                if cover:
                    story.append(PageBreak())
                    cover = False
                story.append(Paragraph(pdf_inline(text),head1 if level == 2 else head2))
        elif kind == "paragraph":
            style = ParagraphStyle("Cover",parent=body,alignment=TA_CENTER,spaceAfter=14) if cover else body
            story.append(Paragraph(pdf_inline(value),style))
        elif kind == "bullet":
            story.append(Paragraph("• "+pdf_inline(value),body))
        elif kind == "number":
            style = ParagraphStyle("ReportNumber",parent=body,spaceAfter=2)
            story.append(Paragraph(pdf_inline(value),style))
        elif kind == "code":
            story.append(Paragraph(escape(value).replace("\n","<br/>"),code_style))
        elif kind == "table":
            count = len(value[0])
            width = document.width
            widths = [width/count]*count
            if count == 5:
                widths = [width*.13]+[width*.2175]*4
            if count == 2:
                widths = [width*.44,width*.56]
            rows = [[Paragraph(pdf_inline(c),cell_style) for c in row] for row in value]
            table = Table(rows,colWidths=widths,repeatRows=1,hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E9F0F5")),
                ("GRID",(0,0),(-1,-1),.4,colors.HexColor("#BCC6CE")),("VALIGN",(0,0),(-1,-1),"TOP"),
                ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
            story.extend([table,Spacer(1,9)])
        elif kind == "image":
            caption,path = value
            image = Image(str(REPORT.parent/path))
            factor = document.width/image.imageWidth
            image.drawWidth,image.drawHeight = document.width,image.imageHeight*factor
            story.append(KeepTogether([image,Paragraph(pdf_inline(caption),caption_style),Spacer(1,8)]))
    def page(canvas,doc):
        canvas.saveState()
        canvas.setFont("BodyFont",8)
        canvas.setFillColor(colors.HexColor("#56616B"))
        canvas.drawCentredString(A4[0]/2,1.05*cm,f"ATBMHTTT | {doc.page}")
        canvas.restoreState()
    document.build(story,onFirstPage=page,onLaterPages=page)


def main():
    if hasattr(sys.stdout,"reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    figures()
    items = list(blocks(REPORT.read_text(encoding="utf-8")))
    word(items)
    pdf(items)
    from pypdf import PdfReader
    pages = PdfReader(REPORT.with_suffix(".pdf")).pages
    print(f"Exported {REPORT.with_suffix('.docx').relative_to(ROOT)}")
    print(f"Exported {REPORT.with_suffix('.pdf').relative_to(ROOT)} ({len(pages)} pages)")


if __name__ == "__main__":
    main()
